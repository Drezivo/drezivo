"""Generate the Drezivo repository-wide onboarding guide from the checked-out scaffold."""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Drezivo-Human-Repository-Architecture-Guide.docx"
SKIP_DIRS = {".git", ".polyrepo-git-archives", "node_modules", "dist", ".next", "coverage", "playwright-report", "test-results"}
REPOS = [
    ("contracts", "Versioned TypeScript/Zod contracts and generated OpenAPI", "api, app, web"),
    ("api", "Express REST API, tenant authorization, Neon/Drizzle persistence, and worker", "app, web"),
    ("app", "Authenticated staff dashboard for rental operations", "staff users"),
    ("web", "Marketing site, public storefront, and guest booking flow", "customers and prospects"),
    ("docs", "Product, architecture, decisions, references, and runbooks", "all workspaces and operators"),
]


def iter_files(repo: Path) -> list[Path]:
    return sorted((p for p in repo.rglob("*") if p.is_file() and not p.name.startswith("~$") and not any(x in SKIP_DIRS for x in p.parts)), key=lambda p: p.as_posix().lower())


def iter_dirs(repo: Path) -> list[Path]:
    return sorted((p for p in repo.rglob("*") if p.is_dir() and not any(x in SKIP_DIRS for x in p.parts)), key=lambda p: p.as_posix().lower())


def para(doc: Document, text: str, style: str | None = None) -> None:
    doc.add_paragraph(text, style=style)


def table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for c, value in zip(t.rows[0].cells, headers): c.text = value
    for row in rows:
        cells = t.add_row().cells
        for c, value in zip(cells, row): c.text = str(value)


def tree(repo: Path, depth: int = 4) -> str:
    lines = [repo.name + "/"]
    def visit(d: Path, prefix: str, level: int) -> None:
        if level > depth: return
        entries = sorted((p for p in d.iterdir() if p.name not in SKIP_DIRS and not p.name.startswith("~$")), key=lambda p: (p.is_file(), p.name.lower()))
        for i, p in enumerate(entries):
            last = i == len(entries) - 1
            lines.append(prefix + ("└── " if last else "├── ") + p.name + ("/" if p.is_dir() else ""))
            if p.is_dir(): visit(p, prefix + ("    " if last else "│   "), level + 1)
    visit(repo, "", 1)
    lines.append("... all deeper files are listed in the file inventory table.")
    return "\n".join(lines)


def role(rel: str, repo: str) -> tuple[str, str]:
    name, low = Path(rel).name, rel.lower()
    if repo == "api" and "src/modules/storefront" in low:
        return {
            "storefront.dto.ts": ("Data Transfer Object (DTO) response types.", "Separates HTTP payloads from database rows; routes map repository results to these shapes."),
            "storefront.repository.ts": ("Storefront Drizzle query boundary.", "Runs parameterized, tenant/branch-scoped queries against Neon."),
            "storefront.routes.ts": ("Express storefront endpoint registration/handlers.", "Composes auth, validation, repository work, response envelopes, and errors."),
            "storefront.schemas.ts": ("Zod input schemas for storefront paths, queries, and bodies.", "Rejects malformed or unknown input before business/database work."),
        }.get(name, ("Storefront module file.", "Participates in the storefront route → validation → repository flow."))
    if name in {"AGENTS.md", "CLAUDE.md", "CLAUDE.local.md.example"}: return ("Agent/developer operating rules.", "Loaded before work; defines security, reliability, testing, and commit behavior.")
    if name == "README.md": return ("Repository entry guide.", "Documents setup, scripts, boundaries, and pointers to canonical specifications.")
    if name == "CHANGELOG.md": return ("Historical release record.", "Preserves prior change context; archived names may remain intentionally.")
    if name == "package.json": return ("Package metadata, scripts, and dependency contract.", "Defines install, lint, test, typecheck, build, and release commands.")
    if name == "package-lock.json": return ("Pinned npm dependency graph.", "Makes installs reproducible and must change with dependency edits.")
    if ".github/workflows" in low: return ("GitHub Actions workflow.", "Runs CI or release automation for this repository.")
    if ".github/" in low: return ("GitHub policy/template/configuration.", "Controls ownership, review, pull requests, and collaboration.")
    if low.endswith(".sql") and "migrations" in low: return ("Ordered database migration.", "Applied to Neon through the API migration process; applied migrations are immutable.")
    if "src/db/schema" in low: return ("Drizzle PostgreSQL table/schema definition.", "Maps TypeScript tables and relations to Neon and query modules.")
    if "src/middleware" in low: return ("Express middleware boundary.", "Handles identity, tenant scope, validation, rate limits, request IDs, or errors before routes.")
    if "src/worker/handlers" in low: return ("Durable background-job handler.", "Claimed by a leased worker with bounded retries and terminal outcomes.")
    if low.endswith("src/worker.ts") or "src/worker/runner" in low: return ("Worker bootstrap/runner.", "Processes outbox or expiry jobs without blocking HTTP requests.")
    if "src/shared" in low: return ("API shared utility.", "Centralizes money, errors, pagination, logging, responses, and result conventions.")
    if "src/common" in low: return ("Shared contract primitive.", "Keeps IDs, money, time, errors, pagination, and idempotency consistent for consumers.")
    if repo == "contracts" and "openapi" in low: return ("OpenAPI generator or REST description.", "Documents the API contract consumed by clients and release checks.")
    if repo == "contracts" and low.startswith("src/"): return ("Versioned contract module.", "Exports schemas/types used by API and both Next.js clients; it has no provider access.")
    if repo in {"app", "web"} and "src/app" in low and low.endswith("page.tsx"): return ("Next.js App Router page.", "Maps a URL to a user journey and calls the shared API client.")
    if repo in {"app", "web"} and "src/app" in low: return ("Next.js route/layout/error/metadata file.", "Defines navigation, shell, health, error, or SEO behavior for its route tree.")
    if repo in {"app", "web"} and "src/components" in low: return ("Reusable React UI component.", "Owns presentation and interaction state; mutations require pending guards.")
    if repo in {"app", "web"} and "src/lib" in low: return ("Frontend library/helper.", "Centralizes API calls, money/status display, permissions, queries, or submit guards.")
    if "tests/" in low or "__tests__" in low or low.endswith((".test.ts", ".test.tsx", ".spec.ts")): return ("Automated test.", "Falsifies a unit, route, UI, or end-to-end behavior before release.")
    if repo == "docs" and "runbooks" in low: return ("Operational runbook.", "Explains deploy, migration, rollback, incident, and recovery actions.")
    if repo == "docs" and "decisions" in low: return ("Architecture Decision Record (ADR).", "Records why a monorepo technical choice exists and its consequences.")
    if repo == "docs" and ("product" in low or "architecture" in low): return ("Canonical product/architecture specification.", "Defines intent and constraints implemented by the workspaces.")
    if name.endswith((".config.ts", ".config.mjs", ".config.js")): return ("Tool configuration.", "Controls compiler, tests, lint, CSS, bundling, or deployment behavior.")
    if name == "Dockerfile": return ("Container build recipe.", "Packages the service with its runtime identity and start command.")
    if name.endswith((".sh", ".mjs")) or (name.endswith(".ts") and rel.startswith("scripts/")): return ("Repository automation script.", "Performs repeatable setup, migration, policy, or checking tasks used by CI/contributors.")
    return ("Source or configuration file.", "Owned by the nearest module; read its local README and agent rules before changing it.")


def inventory(doc: Document, repo: str, path: Path) -> None:
    dirs, files = iter_dirs(path), iter_files(path)
    doc.add_heading("Folder inventory", 2); para(doc, "Each directory is an ownership boundary. The file inventory below is generated from this checkout and explains every file outside dependency/build output.")
    para(doc, tree(path), "No Spacing")
    table(doc, ["Directory", "Purpose", "Connection"], ((p.relative_to(path).as_posix() + "/", role(p.relative_to(path).as_posix(), repo)[0], role(p.relative_to(path).as_posix(), repo)[1]) for p in dirs))
    doc.add_heading("File-by-file inventory", 2)
    table(doc, ["File", "What it does", "How it connects"], ((p.relative_to(path).as_posix(), role(p.relative_to(path).as_posix(), repo)[0], role(p.relative_to(path).as_posix(), repo)[1]) for p in files))


def glossary(doc: Document) -> None:
    doc.add_page_break(); doc.add_heading("Acronyms and deep terms", 1)
    para(doc, "Use these definitions when reading the scaffold. Define unfamiliar terms in the owning document before using them.")
    rows = [
        ("ADR", "Architecture Decision Record", "A dated decision explaining context, choice, and consequences."),
        ("API", "Application Programming Interface", "The Express REST surface clients call, such as creating a reservation hold."),
        ("CI/CD", "Continuous Integration / Continuous Delivery", "Automated checks and controlled release automation."),
        ("DTO", "Data Transfer Object", "A boundary payload shape; it is not automatically a database row."),
        ("ERD", "Entity-Relationship Diagram", "A diagram of data entities/tables and their relationships."),
        ("HMAC", "Hash-based Message Authentication Code", "A keyed digest proving a message was signed by a trusted party."),
        ("HTTP", "Hypertext Transfer Protocol", "Request/response protocol used by browser clients and the API."),
        ("JWT", "JSON Web Token", "Signed identity claims; Clerk verifies identity while the API performs tenant/role authorization."),
        ("NPM", "Node Package Manager", "Installs JavaScript/TypeScript dependencies and runs scripts."),
        ("PII", "Personally Identifiable Information", "Information that identifies a person; keep it out of logs and vault notes."),
        ("PRD", "Product Requirements Document", "Product problem, scope, workflows, acceptance criteria, and release progression."),
        ("REST", "Representational State Transfer", "Resource-oriented HTTP style used between Next.js and Express."),
        ("RLS", "Row-Level Security", "PostgreSQL policies restricting rows; application authorization is still required."),
        ("S3", "Amazon Simple Storage Service", "Object storage for images/evidence; the API issues signed access URLs."),
        ("SaaS", "Software as a Service", "Hosted software shared by isolated business tenants."),
        ("SQL", "Structured Query Language", "Language used by migrations and Neon PostgreSQL queries."),
        ("TRD", "Technical Requirements Document", "Architecture, non-functional requirements, integrations, and constraints."),
        ("UI/UX", "User Interface / User Experience", "Screens and interactions; business truth remains server-side."),
        ("Zod", "TypeScript-first schema validation library", "Validates untrusted input and infers TypeScript types."),
        ("Tenant", "An isolated business account", "One rental company whose staff/branches cannot access another tenant's rows."),
        ("Idempotency key", "Stable key for one mutation intent", "A retry or double-click replays one result instead of duplicating work."),
        ("Outbox", "Durable post-commit work table", "A transaction records work that a worker dispatches after commit."),
        ("Lease", "Time-limited job claim", "A worker owns a job until expiry; another worker can recover it after a crash."),
    ]
    table(doc, ["Term", "Word-for-word meaning", "Plain definition/example"], rows)


def build() -> None:
    doc = Document(); s = doc.sections[0]; s.top_margin = Inches(.65); s.bottom_margin = Inches(.65); s.left_margin = Inches(.7); s.right_margin = Inches(.7)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(9); doc.styles["No Spacing"].font.name = "Consolas"; doc.styles["No Spacing"].font.size = Pt(7)
    doc.add_heading("Drezivo Monorepo Architecture and Onboarding Guide", 0)
    p = doc.add_paragraph("Root-wide human orientation for engineers and operators"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, f"Generated from the checked-out scaffold on {date.today().isoformat()}.")
    para(doc, "Purpose. This is the pre-development map: it explains every root-level item, every tracked workspace file and directory, ownership boundaries, file-to-file connections, release order, acronym definitions, and how to join the shared Obsidian knowledge base.")
    para(doc, "Current maturity. The repositories are an incomplete scaffold. A recognized but unfinished API route may return HTTP 501; no production Neon database or provider wiring is implied by the presence of schemas/configuration.")
    doc.add_heading("How to use this guide", 1)
    for x in ["Start with the root tree and system connection map.", "Read the repository section for the area you will change; the file inventory explains why each file exists.", "Read the nearest AGENTS.md and CLAUDE.md before editing.", "Trace a feature in dependency order: contracts → api → app/web → docs/runbooks.", "Treat the PRD, TRD, data model, accepted ADRs, and migrations as sources of truth.", "Use the final onboarding section to open docs/second-brain and keep durable decisions linked."]: para(doc, "☐ " + x)
    doc.add_heading("System at a glance", 1)
    para(doc, "Connecting tree\nDrezivo monorepo\n├── contracts  → shared workspace package for schemas, enums, money, errors, OpenAPI\n├── api        → Express REST service, Clerk authz, Drizzle/Neon, worker\n├── app        → authenticated staff dashboard\n├── web        → marketing, public storefront, guest booking\n└── docs       → PRD, TRD, ERD, ADRs, screenshots, runbooks")
    para(doc, "Runtime flow\nBrowser (app/web) → Clerk session → API middleware → Zod validation → tenant + branch authorization → Drizzle transaction → Neon PostgreSQL\n                                                                                     ├→ S3 signed object access\n                                                                                     └→ outbox/lease worker for durable side effects")
    table(doc, ["Repository", "Owns", "Must never own"], [("contracts", "Stable request/response types, enums, OpenAPI", "DB connections, secrets, UI, authorization decisions"), ("api", "Rules, authz, transactions, persistence, jobs", "Client totals/roles or fire-and-forget effects"), ("app", "Staff screens and interaction states", "Direct Neon access or authoritative money/roles"), ("web", "Public/guest screens, marketing, SEO", "Private staff data or server authz"), ("docs", "Intent, decisions, operations, evidence", "Drifting implementation copies")])
    doc.add_heading("Root workspace inventory", 1); para(doc, "The root is one Git repository containing five workspaces. Root documents guide the applications, service, package, and operators.")
    root_roles = {
        "package.json": "Root workspace scripts and dependency policy.",
        "package-lock.json": "One reproducible dependency graph for the monorepo.",
        "LICENSE.md": "Proprietary license for the organization repository.",
        "SECURITY-FOUNDATION.template.md": "Pointer to the private security baseline.",
        "AGENTS.md": "Root agent/contributor operating contract.",
        "CLAUDE.md": "Root Claude entry point.",
        "README.md": "Root reading order and scaffold links.",
        "ROOT-REPOSITORY-ARCHITECTURE.md": "Markdown source-of-truth architecture map.",
        "AI-AGENT-ONBOARDING.md": "Machine-facing first-turn workflow for agents.",
        "MEMORY_BANK.md": "Durable project lessons.",
        "Drezivo-Human-Repository-Architecture-Guide.docx": "Human visual onboarding guide.",
        ".claude": "Claude rules and skills.",
        ".codex": "Codex rules and skills mirror.",
        "docs": "Canonical specifications, runbooks, and the project-only Obsidian vault.",
        "graphify-out": "Derived graph artifacts from original PRD.",
        "RootResource": "Supplied design and reference assets.",
        "dbml-error.log": "Historical validation log.",
        "qa-docx": "Document QA scratch output, if present.",
        "build_arch_guide.py": "Reproducible DOCX generator.",
    }
    entries = sorted((p for p in ROOT.iterdir() if not p.name.startswith("~$")), key=lambda p: (p.is_file(), p.name.lower()))
    table(doc, ["Root item", "What belongs here", "Connection"], ((e.name + ("/" if e.is_dir() else ""), root_roles.get(e.name, "Workspace item; inspect owner before changing."), "Shared onboarding/configuration." if e.name not in {x[0] for x in REPOS} else "Workspace boundary; see its section.") for e in entries))
    para(doc, "Top-level tree\n" + "\n".join("├── " + e.name + ("/" if e.is_dir() else "") for e in entries))
    for repo, purpose, consumers in REPOS:
        doc.add_page_break(); doc.add_heading(f"Workspace: {repo}", 1); para(doc, f"Purpose: {purpose}. Primary relationship: {consumers}.")
        para(doc, "Change rule. Keep edits inside this ownership boundary. If a request/response changes, update contracts first; if persistence changes, update API migrations/schema and the data model/runbook; then update consumers and tests.")
        inventory(doc, repo, ROOT / repo)
        if repo == "api":
            doc.add_heading("API storefront module deep dive", 2); para(doc, "The src/modules/storefront directory is split so each boundary has one job. DTO, schema, repository, and route files prevent database rows, unvalidated input, and HTTP orchestration from becoming one tangled module.")
            table(doc, ["File", "Function", "Sequence"], [("storefront.schemas.ts", "Rejects unknown/invalid path, query, and body values.", "middleware → parse → handler"), ("storefront.dto.ts", "Defines response payloads separate from rows.", "repository → mapper → envelope"), ("storefront.repository.ts", "Runs parameterized, scoped Drizzle queries.", "route → repository → Neon"), ("storefront.routes.ts", "Registers endpoints and composes boundaries.", "HTTP → auth/tenant → route → JSON"), ("tests / __tests__", "Falsifies behavior and prevents regressions.", "CI → assertions → release evidence")])
    glossary(doc)
    doc.add_page_break(); doc.add_heading("Monorepo development and release", 1)
    table(doc, ["Order", "Action", "Evidence"], [("1 Contract", "Update schema/enum/error/money/OpenAPI.", "Contract tests and reviewed compatibility diff."), ("2 API", "Implement authz, transaction, persistence, idempotency, outbox.", "Typecheck/lint/tests plus sequential/concurrent double-fire test."), ("3 App/web", "Wire the correct user journey.", "Pending disabled state, guard, errors/empty state, UI tests."), ("4 Docs", "Update PRD/TRD/data model/ADR/runbook and Obsidian note.", "Links resolve and decision is recorded."), ("5 Release", "CI, staging smoke, migration, monitoring, rollback.", "Attached command output and known limitations.")])
    para(doc, "Commit format: Conventional Commits, for example feat(api): add storefront availability, fix(web): guard duplicate hold submission, or docs: update migration rollback. Keep one coherent change per commit and explain breaking contract changes.")
    doc.add_heading("Joining the shared Obsidian second brain", 1); para(doc, "Vault path: docs/second-brain/. Start at 00-Home/Drezivo Home.md. The vault is a project-only thinking/navigation layer; code, migrations, PRD/TRD, and accepted ADRs remain authoritative.")
    para(doc, "Vault tree\ndocs/second-brain/\n├── 00-Home          → navigation hub\n├── 01-Product       → product brief and scope links\n├── 02-Architecture  → architecture note and JSON Canvas graph\n├── 03-Repositories  → ownership/file map\n├── 04-Decisions     → decision register and vault policy\n├── 05-Operations    → hosting, GitHub rename, operating model\n├── 06-Research      → source/date/relevance register\n├── 07-Glossary      → shared terms\n├── 08-Daily         → transient capture promoted to durable notes\n├── 99-Archive       → superseded notes with status: archived\n└── _attachments     → reviewed, non-secret assets")
    table(doc, ["Step", "Action", "Safety rule"], [("Open", "Open the folder as an Obsidian vault; start from Drezivo Home.", "Use this project vault only."), ("Share", "Private Obsidian Sync with end-to-end encryption, or private Git.", "Invite named collaborators only."), ("Write", "Use YAML properties, vault-relative wikilinks, source/date fields.", "No tokens, Neon URLs, payment evidence, or PII."), ("Graph", "Update the Canvas when relationships change; keep node IDs stable.", "Validate JSON and targets before sync."), ("Promote", "Move durable conclusions from daily notes into permanent notes.", "Record decisions; archive instead of erasing history.")])
    para(doc, "Project-local .codex/skills/obsidian-skills/ and .claude/skills/obsidian-skills/ follow the upstream Agent Skills package. Use obsidian-markdown for notes, json-canvas for graph files, obsidian-bases for metadata views, and obsidian-cli only with a running Obsidian desktop app.")
    doc.add_heading("First-day verification checklist", 1)
    for x in ["Confirm workspace tree and active branch.", "Read root and target repository agent rules/README.", "Install dependencies with the lockfile; never commit node_modules or secrets.", "Trace one read and one mutation UI → contract → API → database/outbox.", "Verify every mutation has pending disabled state, early-return guard, and one idempotency key per intent.", "Run typecheck, lint, tests, and build applicable to the change.", "Update owning docs and Obsidian note with sources/dates.", "Review diff for scope leaks, PII/secrets, stale names, and generated output."]: para(doc, "☐ " + x)
    footer = doc.sections[0].footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("Drezivo Repository Architecture Guide | generated from scaffold inventory").font.size = Pt(8)
    doc.save(OUT); print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__": build()

